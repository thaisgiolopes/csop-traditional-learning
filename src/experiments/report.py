import csv
from dataclasses import asdict
from datetime import datetime
import json
import os
from pathlib import Path
from typing import Any, Iterable
from .database import (
    ExperimentDatabase,
    ExperimentRecord,
)
from openpyxl import Workbook, load_workbook


SUMMARY_COLUMNS = (
    "experiment_id",
    "name",
    "description",
    "status",
    "created_at",
    "started_at",
    "finished_at",
    "total_duration_seconds",
    "graph_reference_json",
    "subgraph_generation_method",
    "sample_count",
    "path_length",
    "feature_set_id",
    "feature_names_json",
    "number_of_features",
    "dataset_split_strategy",
    "train_ratio",
    "test_ratio",
    "model_type",
    "model_hyperparameters_json",
    "recorded_stage_count",
    "prediction_count",
    "evaluation_metrics_json",
)

XLSX_SHEETS = {
    "Experiments": (
        "experiment_id",
        "experiment_name",
        "description",
        "status",
        "created_at",
        "started_at",
        "finished_at",
        "graph_id",
        "graph_path",
        "graph_num_vertices",
        "graph_num_edges",
        "generation_method",
        "sample_count",
        "path_length",
        "generation_seed",
        "feature_set_id",
        "feature_names",
        "num_features",
        "objective_name",
        "split_strategy",
        "train_ratio",
        "test_ratio",
        "split_seed",
        "model_type",
        "model_hyperparameters",
        "total_duration_seconds",
        "peak_or_recorded_memory_bytes",
        "num_predictions",
        "evaluation_metrics",
        "configuration_path",
        "environment_path",
    ),
    "Stages": (
        "experiment_id",
        "stage_id",
        "stage_name",
        "status",
        "started_at",
        "finished_at",
        "duration_seconds",
        "memory_before_bytes",
        "memory_after_bytes",
        "memory_delta_bytes",
    ),
    "Artifacts": (
        "artifact_id",
        "experiment_id",
        "category",
        "name",
        "artifact_type",
        "path",
        "size_bytes",
        "created_at",
    ),
    "Evaluations": (
        "experiment_id",
        "evaluation_id",
        "metric_name",
        "metric_value",
    ),
    "Predictions": (
        "experiment_id",
        "prediction_id",
        "sample_id",
        "predicted_value",
        "actual_value",
        "prediction_error",
    ),
}


def _json_cell(value: Any) -> str:
    """Serialize complex values deterministically for a CSV cell."""
    return json.dumps(
        value,
        ensure_ascii=False,
        sort_keys=True,
        separators=(",", ":"),
        allow_nan=False,
    )


def _total_duration_seconds(
    started_at: str | None,
    finished_at: str | None,
) -> float | None:
    """Calculate duration from recorded ISO-8601 timestamps."""
    if not started_at or not finished_at:
        return None

    try:
        start = datetime.fromisoformat(started_at.replace("Z", "+00:00"))
        finish = datetime.fromisoformat(finished_at.replace("Z", "+00:00"))
    except ValueError:
        return None

    return (finish - start).total_seconds()


class ExperimentReport:
    """Generate tabular reports from records in the experiment database."""

    def __init__(
        self, 
        database: ExperimentDatabase,
        project_root: str | Path | None = None,
    ) -> None:
        self._database = database
        self._project_root = (
            Path(project_root).resolve()
            if project_root is not None
            else Path(__file__).resolve().parents[2]
        )

    def _project_relative_path(self, value: str | Path | None) -> str:
        """Return a filesystem path relative to the project root."""
        if value is None or str(value).strip() == "":
            return ""

        path = Path(value).expanduser()
        if not path.is_absolute():
            return path.as_posix()

        return Path(
            os.path.relpath(path.resolve(), start=self._project_root)
        ).as_posix()

    @staticmethod
    def _read_graph_metadata(
        graph_path: str | None,
    ) -> tuple[int | str, int | str]:
        """Read vertex and edge counts from the configured graph metadata file."""
        if not graph_path or "://" in graph_path:
            return "", ""

        metadata_path = Path(graph_path).expanduser() / "metadata"

        try:
            values = metadata_path.read_text(
                encoding="utf-8"
            ).split()
            if len(values) < 2:
                return "", ""

            return int(values[0]), int(values[1])
        except (OSError, ValueError):
            return "", ""
        
    def write_csv(
        self,
        path: str | Path,
        experiment_ids: str | Iterable[str] | None = None,
    ) -> Path:
        """Write one summary row per selected experiment to UTF-8 CSV."""
        rows = self._summary_rows(experiment_ids)
        output_path = Path(path)
        output_path.parent.mkdir(parents=True, exist_ok=True)

        with output_path.open("w", encoding="utf-8", newline="") as csv_file:
            writer = csv.DictWriter(
                csv_file,
                fieldnames=SUMMARY_COLUMNS,
                extrasaction="raise",
            )
            writer.writeheader()
            writer.writerows(rows)

        return output_path

    def update_xlsx(
        self,
        path: str | Path,
        experiment_id: str,
    ) -> Path:
        """Insert or update only one experiment's rows in the workbook."""
        output_path = Path(path)
        if output_path.suffix.lower() != ".xlsx":
            raise ValueError("Report path must have an .xlsx extension.")

        output_path.parent.mkdir(parents=True, exist_ok=True)

        experiment = self._database.get_experiment(experiment_id)
        config = self._load_configuration(experiment.configuration_path)

        stages = self._database.get_experiment_stages(experiment_id)
        artifacts = self._database.get_experiment_artifacts(experiment_id)
        predictions = self._database.get_experiment_predictions(experiment_id)
        evaluations = self._database.get_experiment_evaluations(experiment_id)

        summary_row = self._xlsx_experiment_row(
            experiment,
            config,
            stages,
            predictions,
            evaluations,
        )
        stage_rows = [
            {
                "experiment_id": item.experiment_id,
                "stage_id": item.stage_id,
                "stage_name": item.stage_name,
                "status": item.status,
                "started_at": item.started_at,
                "finished_at": item.finished_at,
                "duration_seconds": item.duration_seconds,
                "memory_before_bytes": item.memory_before_bytes,
                "memory_after_bytes": item.memory_after_bytes,
                "memory_delta_bytes": item.memory_delta_bytes,
            }
            for item in stages
        ]
        artifact_rows = [
            {
                "artifact_id": item.artifact_id,
                "experiment_id": item.experiment_id,
                "category": item.category,
                "name": item.name,
                "artifact_type": item.artifact_type,
                "path": self._project_relative_path(item.path),
                "size_bytes": item.size_bytes,
                "created_at": item.created_at,
            }
            for item in artifacts
        ]
        evaluation_rows = [
            {
                "experiment_id": item.experiment_id,
                "evaluation_id": item.evaluation_id,
                "metric_name": item.metric_name,
                "metric_value": item.metric_value,
            }
            for item in evaluations
        ]
        prediction_rows = [
            {
                "experiment_id": item.experiment_id,
                "prediction_id": item.prediction_id,
                "sample_id": item.sample_id,
                "predicted_value": item.predicted_value,
                "actual_value": item.actual_value,
                "prediction_error": item.prediction_error,
            }
            for item in predictions
        ]

        if output_path.exists():
            workbook = load_workbook(output_path)
        else:
            workbook = Workbook()
            workbook.remove(workbook.active)

        worksheets = {
            name: self._get_or_create_sheet(workbook, name, headers)
            for name, headers in XLSX_SHEETS.items()
        }

        self._upsert_summary(
            worksheets["Experiments"],
            XLSX_SHEETS["Experiments"],
            summary_row,
        )
        self._replace_experiment_rows(
            worksheets["Stages"],
            XLSX_SHEETS["Stages"],
            experiment_id,
            stage_rows,
        )
        self._replace_experiment_rows(
            worksheets["Artifacts"],
            XLSX_SHEETS["Artifacts"],
            experiment_id,
            artifact_rows,
        )
        self._replace_experiment_rows(
            worksheets["Evaluations"],
            XLSX_SHEETS["Evaluations"],
            experiment_id,
            evaluation_rows,
        )
        self._replace_experiment_rows(
            worksheets["Predictions"],
            XLSX_SHEETS["Predictions"],
            experiment_id,
            prediction_rows,
        )

        for worksheet in worksheets.values():
            worksheet.freeze_panes = "A2"
            worksheet.auto_filter.ref = worksheet.dimensions

        temporary_path = output_path.with_name(
            f".{output_path.stem}.tmp{output_path.suffix}"
        )
        try:
            workbook.save(temporary_path)
            os.replace(temporary_path, output_path)
        finally:
            if temporary_path.exists():
                temporary_path.unlink()

        return output_path

    def _xlsx_experiment_row(
        self,
        experiment: ExperimentRecord,
        config: dict[str, Any],
        stages: list[Any],
        predictions: list[Any],
        evaluations: list[Any],
    ) -> dict[str, Any]:
        graph = config.get("graph") or {}
        graph_parameters = graph.get("parameters") or {}
        graph_source_path = graph.get("path") or graph.get("source")
        generation = config.get("subgraph_generation") or {}
        features = config.get("features") or {}
        split = config.get("dataset_split") or {}
        model = config.get("model") or {}
        feature_names = features.get("feature_names") or []

        graph_num_vertices, graph_num_edges = self._read_graph_metadata(
            graph.get("path")
        )

        memory_values = [
            value
            for stage in stages
            for value in (
                stage.memory_before_bytes,
                stage.memory_after_bytes,
            )
            if value is not None
        ]
        metrics = {
            item.metric_name: item.metric_value
            for item in sorted(evaluations, key=lambda item: item.metric_name)
        }

        return {
            "experiment_id": experiment.experiment_id,
            "experiment_name": experiment.name or "",
            "description": experiment.description or "",
            "status": experiment.status,
            "created_at": experiment.created_at,
            "started_at": experiment.started_at or "",
            "finished_at": experiment.finished_at or "",
            "graph_id": graph.get("graph_id") or "",
            "graph_path": self._project_relative_path(graph_source_path),
            "graph_num_vertices": (
                graph_num_vertices
                if graph_num_vertices != ""
                else graph_parameters.get("num_vertices", "")
            ),
            "graph_num_edges": (
                graph_num_edges
                if graph_num_edges != ""
                else graph_parameters.get("num_edges", "")
            ),
            "generation_method": generation.get("method") or "",
            "sample_count": generation.get("sample_count", ""),
            "path_length": (
                generation.get("path_length")
                if generation.get("path_length") is not None
                else ""
            ),
            "generation_seed": (
                generation.get("seed")
                if generation.get("seed") is not None
                else ""
            ),
            "feature_set_id": features.get("feature_set_id") or "",
            "feature_names": _json_cell(feature_names),
            "num_features": len(feature_names),
            "objective_name": features.get("objective_name") or "",
            "split_strategy": split.get("strategy") or "",
            "train_ratio": split.get("train_ratio", ""),
            "test_ratio": split.get("test_ratio", ""),
            "split_seed": (
                split.get("seed")
                if split.get("seed") is not None
                else ""
            ),
            "model_type": model.get("model_type") or "",
            "model_hyperparameters": _json_cell(
                model.get("hyperparameters") or {}
            ),
            "total_duration_seconds": _total_duration_seconds(
                experiment.started_at,
                experiment.finished_at,
            ),
            "peak_or_recorded_memory_bytes": (
                max(memory_values) if memory_values else ""
            ),
            "num_predictions": len(predictions),
            "evaluation_metrics": _json_cell(metrics),
            "configuration_path": self._project_relative_path(
                experiment.configuration_path
            ),
            "environment_path": self._project_relative_path(
                experiment.environment_path
            ),
        }

    @staticmethod
    def _load_json_artifact(path: str | None) -> dict[str, Any]:
        """Load JSON only from a path explicitly registered in SQLite."""
        if not path:
            return {}

        try:
            with Path(path).open("r", encoding="utf-8") as json_file:
                value = json.load(json_file)
        except (OSError, json.JSONDecodeError):
            return {}

        return value if isinstance(value, dict) else {}

    @staticmethod
    def _get_or_create_sheet(workbook, name, headers):
        if name in workbook.sheetnames:
            worksheet = workbook[name]
            existing_headers = tuple(
                cell.value for cell in worksheet[1]
            )
            if existing_headers != headers:
                raise ValueError(
                    f"Unexpected columns in XLSX sheet {name!r}."
                )
            return worksheet

        worksheet = workbook.create_sheet(name)
        worksheet.append(headers)
        return worksheet

    @staticmethod
    def _upsert_summary(worksheet, headers, row):
        experiment_id_column = headers.index("experiment_id") + 1

        for row_number in range(2, worksheet.max_row + 1):
            if worksheet.cell(row_number, experiment_id_column).value == row[
                "experiment_id"
            ]:
                for column_number, header in enumerate(headers, start=1):
                    worksheet.cell(row_number, column_number).value = row[header]
                return

        worksheet.append([row[header] for header in headers])

    @staticmethod
    def _replace_experiment_rows(
        worksheet,
        headers,
        experiment_id,
        rows,
    ):
        experiment_id_column = headers.index("experiment_id") + 1

        for row_number in range(worksheet.max_row, 1, -1):
            if worksheet.cell(row_number, experiment_id_column).value == experiment_id:
                worksheet.delete_rows(row_number)

        for row in rows:
            worksheet.append([row[header] for header in headers])

    def get_experiment_details(self, experiment_id: str) -> dict[str, Any]:
        """Return a summary and database records for one experiment."""
        experiment = self._database.get_experiment(experiment_id)
        summary = self._summary_row(experiment)

        return {
            "summary": summary,
            "stages": [
                asdict(record)
                for record in self._database.get_experiment_stages(
                    experiment_id
                )
            ],
            "artifacts": [
                asdict(record)
                for record in self._database.get_experiment_artifacts(
                    experiment_id
                )
            ],
            "predictions": [
                asdict(record)
                for record in self._database.get_experiment_predictions(
                    experiment_id
                )
            ],
            "evaluations": [
                asdict(record)
                for record in self._database.get_experiment_evaluations(
                    experiment_id
                )
            ],
        }

    def _summary_rows(
        self,
        experiment_ids: str | Iterable[str] | None,
    ) -> list[dict[str, Any]]:
        if experiment_ids is None:
            experiments = self._database.list_experiments()
        else:
            requested_ids = (
                [experiment_ids]
                if isinstance(experiment_ids, str)
                else list(dict.fromkeys(experiment_ids))
            )
            experiments = [
                self._database.get_experiment(experiment_id)
                for experiment_id in requested_ids
            ]

        return [self._summary_row(item) for item in experiments]

    def _summary_row(
        self,
        experiment: ExperimentRecord,
    ) -> dict[str, Any]:
        config = self._load_configuration(experiment.configuration_path)
        graph = config.get("graph") or {}
        generation = config.get("subgraph_generation") or {}
        features = config.get("features") or {}
        split = config.get("dataset_split") or {}
        model = config.get("model") or {}

        stages = self._database.get_experiment_stages(
            experiment.experiment_id
        )
        predictions = self._database.get_experiment_predictions(
            experiment.experiment_id
        )
        evaluations = self._database.get_experiment_evaluations(
            experiment.experiment_id
        )

        graph_reference = {
            key: graph[key]
            for key in ("graph_id", "path", "source")
            if graph.get(key) is not None
        }
        feature_names = features.get("feature_names") or []
        evaluation_metrics = {
            record.metric_name: record.metric_value
            for record in sorted(
                evaluations,
                key=lambda item: item.metric_name,
            )
        }

        return {
            "experiment_id": experiment.experiment_id,
            "name": experiment.name or "",
            "description": experiment.description or "",
            "status": experiment.status,
            "created_at": experiment.created_at,
            "started_at": experiment.started_at or "",
            "finished_at": experiment.finished_at or "",
            "total_duration_seconds": _total_duration_seconds(
                experiment.started_at,
                experiment.finished_at,
            ),
            "graph_reference_json": _json_cell(graph_reference),
            "subgraph_generation_method": generation.get("method", ""),
            "sample_count": generation.get("sample_count", ""),
            "path_length": (
                generation.get("path_length")
                if generation.get("path_length") is not None
                else ""
            ),
            "feature_set_id": features.get("feature_set_id", ""),
            "feature_names_json": _json_cell(feature_names),
            "number_of_features": len(feature_names),
            "dataset_split_strategy": split.get("strategy", ""),
            "train_ratio": split.get("train_ratio", ""),
            "test_ratio": split.get("test_ratio", ""),
            "model_type": model.get("model_type", ""),
            "model_hyperparameters_json": _json_cell(
                model.get("hyperparameters") or {}
            ),
            "recorded_stage_count": len(stages),
            "prediction_count": len(predictions),
            "evaluation_metrics_json": _json_cell(evaluation_metrics),
        }

    @staticmethod
    def _load_configuration(path: str | None) -> dict[str, Any]:
        """Load configuration from the exact artifact path stored in SQLite."""
        if not path:
            return {}

        try:
            with Path(path).open("r", encoding="utf-8") as json_file:
                configuration = json.load(json_file)
        except (OSError, json.JSONDecodeError):
            return {}

        return configuration if isinstance(configuration, dict) else {}
"""End-to-end smoke coverage for the configured experiment runner."""

import csv
import json
from pathlib import Path

import networkx as nx
import numpy as np
import pytest
from openpyxl import load_workbook

from src.application.csop_experiment_runner import run_experiment
from src.application.csop_experiment_runner import run_experiment_batch
from src.experiments.config import (
    DatasetSplitConfig,
    ExperimentConfig,
    FeatureConfig,
    GraphConfig,
    ModelConfig,
    SubgraphGenerationConfig,
)
from src.experiments.database import ExperimentDatabase
from src.graph import Graph
from src.graph.loader import GraphLoader
from src.model.lightgbm_model import LightGBMModel
from src.objectives.tds import TDSObjective


SEED = 29
SAMPLE_COUNT = 24
SUBGRAPH_SIZE = 4
FEATURE_NAMES = [
    "num_vertices",
    "num_edges",
    "degree_mean",
    "degree_max",
    "degree_min",
    "degree_std",
]


def _write_graph(path: Path) -> nx.Graph:
    """Write a reproducible, small graph in the GraphLoader input format."""
    graph = nx.gnp_random_graph(10, 0.42, seed=SEED)
    assert nx.is_connected(graph)
    path.mkdir(parents=True)
    (path / "metadata").write_text(
        f"{graph.number_of_nodes()} {graph.number_of_edges()}\n",
        encoding="utf-8",
    )
    edge_ids = {
        frozenset(edge): index
        for index, edge in enumerate(sorted(graph.edges()))
    }
    lines = []
    for vertex in graph.nodes:
        entries = [
            f"{neighbor},{edge_ids[frozenset((vertex, neighbor))]}"
            for neighbor in sorted(graph.neighbors(vertex))
        ]
        lines.append(" ".join(entries))
    (path / "adjlists").write_text("\n".join(lines) + "\n", encoding="utf-8")
    return graph


def _config(graph_path: Path, method: str = "networkx_connected_expansion"):
    generation = SubgraphGenerationConfig(
        method=method,
        sample_count=SAMPLE_COUNT,
        path_length=3 if method == "networkx_random_paths" else None,
        seed=SEED,
        parameters=(
            {"subgraph_size": SUBGRAPH_SIZE, "max_attempts": 5000}
            if method == "networkx_connected_expansion"
            else {}
        ),
    )
    return ExperimentConfig(
        name=f"smoke-{method}",
        description="Small deterministic runner integration smoke test.",
        random_seed=SEED,
        graph=GraphConfig(graph_id="smoke_graph", path=graph_path),
        subgraph_generation=generation,
        features=FeatureConfig(
            feature_set_id="graph-size-and-local-degree-v1",
            feature_names=FEATURE_NAMES,
            objective_name="tds",
        ),
        dataset_split=DatasetSplitConfig(
            strategy="random",
            test_ratio=0.2,
            validation_size=0.25,
            random_state=SEED,
        ),
        model=ModelConfig(
            model_type="lightgbm",
            hyperparameters={
                "n_estimators": 8,
                "learning_rate": 0.1,
                "num_leaves": 5,
                "min_child_samples": 1,
                "verbosity": -1,
                "random_state": SEED,
                "n_jobs": 1,
            },
        ),
    )


def _run(config, root: Path) -> str:
    return run_experiment(
        config,
        experiments_root=root / "runs",
        database_path=root / "experiments.sqlite3",
        report_path=root / "reports" / "experiments.xlsx",
        processed_root=root / "processed",
        model_factory=lambda model_config: LightGBMModel(
            **model_config.hyperparameters
        ),
        model_serializer=lambda model: model._estimator.booster_.model_to_string(),
    )


def _experiment_root(root: Path, experiment_id: str) -> Path:
    return root / "runs" / experiment_id


def test_configured_runner_end_to_end_smoke_and_reproducibility(tmp_path):
    graph_path = tmp_path / "graph"
    source_nx = _write_graph(graph_path)
    source_nx = GraphLoader(graph_path).load().networkx_graph
    root = tmp_path / "connected-expansion"
    config = _config(graph_path)

    first_id = _run(config, root)
    second_id = _run(config, root)
    first_root = _experiment_root(root, first_id)
    second_root = _experiment_root(root, second_id)

    with ExperimentDatabase(root / "experiments.sqlite3") as database:
        for experiment_id, experiment_root in (
            (first_id, first_root),
            (second_id, second_root),
        ):
            experiment = database.get_experiment(experiment_id)
            assert experiment.status == "COMPLETED"

            request = database.get_experiment_sample_request(experiment_id)
            assert request.requested_count == SAMPLE_COUNT
            assert request.total_available == SAMPLE_COUNT

            associations = database.get_experiment_samples(experiment_id)
            assert len(associations) == SAMPLE_COUNT
            assert len({item.sample_id for item in associations}) == SAMPLE_COUNT
            sample_files = list((experiment_root / "subgraphs" / "smoke_graph").glob("*.json"))
            assert len(sample_files) == SAMPLE_COUNT
            vertex_sets = []
            objective = TDSObjective()
            for sample_file in sample_files:
                sample = json.loads(sample_file.read_text(encoding="utf-8"))
                vertices = set(sample["vertices"])
                vertex_sets.append(frozenset(vertices))
                candidate = Graph(vertices, sample["edges"])
                assert len(vertices) == SUBGRAPH_SIZE
                assert nx.is_connected(candidate.networkx_graph)
                assert vertices <= set(source_nx.nodes)
                expected = source_nx.subgraph(vertices)
                assert set(map(frozenset, candidate.edges)) == set(
                    map(frozenset, expected.edges)
                )
                assert objective.compute(candidate) == pytest.approx(
                    sum(nx.triangles(expected).values()) / (3 * SUBGRAPH_SIZE)
                )
            assert len(set(vertex_sets)) == SAMPLE_COUNT

            all_samples_path = experiment_root / "datasets" / "dataset.csv"
            with all_samples_path.open(newline="", encoding="utf-8") as handle:
                rows = list(csv.DictReader(handle))
            assert len(rows) == SAMPLE_COUNT
            assert len(rows[0]) - 4 == len(FEATURE_NAMES)
            assert all(np.isfinite(float(row["target"])) for row in rows)

            split_paths = {
                name: experiment_root / "datasets" / f"{name}_dataset.csv"
                for name in ("train", "validation", "test")
            }
            split_rows = {}
            for name, path in split_paths.items():
                with path.open(newline="", encoding="utf-8") as handle:
                    split_rows[name] = list(csv.DictReader(handle))
                assert split_rows[name]
                assert len(split_rows[name][0]) - 4 == len(FEATURE_NAMES)
            split_ids = {
                name: {row["sample_id"] for row in values}
                for name, values in split_rows.items()
            }
            assert split_ids["train"].isdisjoint(split_ids["validation"])
            assert split_ids["train"].isdisjoint(split_ids["test"])
            assert split_ids["validation"].isdisjoint(split_ids["test"])
            assert set.union(*split_ids.values()) == {
                row["sample_id"] for row in rows
            }

            metrics = json.loads(
                (experiment_root / "evaluation" / "metrics.json").read_text(
                    encoding="utf-8"
                )
            )
            assert set(metrics) == {"train", "validation", "test"}
            for result in metrics.values():
                assert result["num_samples"] > 0
                assert set(result["metrics"]) == {"mae", "rmse", "r2"}
                assert all(np.isfinite(value) for value in result["metrics"].values())

            stages = database.get_experiment_stages(experiment_id)
            assert stages
            assert all(stage.status == "completed" for stage in stages)
            assert all(stage.duration_seconds is not None for stage in stages)
            assert all(stage.memory_before_bytes is not None for stage in stages)
            assert all(stage.memory_after_bytes is not None for stage in stages)
            assert {stage.stage_name for stage in stages} >= {
                "candidate_generation",
                    "feature_extraction_and_targets",
                "train_validation_test_split",
                "model_training",
                "test_evaluation",
                "report_persistence",
            }
            persisted_metrics = database.get_experiment_metrics(experiment_id)
            assert len(persisted_metrics) == 9
            for metric in persisted_metrics:
                assert metric.metric_value == pytest.approx(
                    metrics[metric.split]["metrics"][metric.metric_name]
                )

        first_predictions = (first_root / "predictions" / "predictions.csv").read_bytes()
        second_predictions = (second_root / "predictions" / "predictions.csv").read_bytes()
        assert first_predictions == second_predictions

        first_request = database.get_experiment_sample_request(first_id)
        second_request = database.get_experiment_sample_request(second_id)
        assert first_request.generated_count == SAMPLE_COUNT
        assert second_request.generated_count == 0
        assert second_request.reused_count == SAMPLE_COUNT

    report_path = root / "reports" / "experiments.xlsx"
    assert report_path.is_file()
    workbook = load_workbook(report_path, read_only=True, data_only=True)
    try:
        report_ids = {
            row[0] for row in list(workbook["Experiments"].values)[1:]
        }
    finally:
        workbook.close()
    assert {first_id, second_id} <= report_ids

    # Preserve coverage of the alternate strategy through the public runner.
    random_root = tmp_path / "random-paths"
    random_id = _run(_config(graph_path, "networkx_random_paths"), random_root)
    with ExperimentDatabase(random_root / "experiments.sqlite3") as database:
        assert database.get_experiment(random_id).status == "COMPLETED"
        random_request = database.get_experiment_sample_request(random_id)
        assert random_request.requested_count == SAMPLE_COUNT
        assert random_request.generated_count == SAMPLE_COUNT


def test_grouped_sizes_train_one_shared_model_and_persist_size_metrics(tmp_path):
    graph_path = tmp_path / "grouped-graph"
    _write_graph(graph_path)
    root = tmp_path / "grouped"
    configs = []
    for size in (3, 4):
        config = _config(graph_path)
        config.name = f"grouped-k{size}"
        config.graph.graph_id = f"smoke_k{size}"
        config.subgraph_generation.parameters["subgraph_size"] = size
        config.features.feature_set_id = (
            "graph-size-local-degree-and-subgraph-size-v2"
        )
        config.features.feature_names = [
            *FEATURE_NAMES[:2],
            "subgraph_size",
            *FEATURE_NAMES[2:],
        ]
        config.output = {
            "experiments_root": str(root / "runs"),
            "database_path": str(root / "experiments.sqlite3"),
            "report_path": str(root / "reports" / "experiments.xlsx"),
            "processed_root": str(root / "processed"),
        }
        configs.append(config)

    model_factory_calls = []

    def model_factory(model_config):
        model_factory_calls.append(model_config.model_type)
        return LightGBMModel(**model_config.hyperparameters)

    experiment_id = run_experiment_batch(
        configs,
        model_factory=model_factory,
        model_serializer=lambda model: model._estimator.booster_.model_to_string(),
    )
    experiment_root = root / "runs" / experiment_id

    with ExperimentDatabase(root / "experiments.sqlite3") as database:
        assert database.get_experiment(experiment_id).status == "COMPLETED"
        assert len(database.list_experiments()) == 1
        requests = database.get_experiment_sample_requests(experiment_id)
        assert len(requests) == 2
        assert {request.requested_count for request in requests} == {SAMPLE_COUNT}
        assert len(database.get_experiment_samples(experiment_id)) == 2 * SAMPLE_COUNT

        metrics = database.get_experiment_metrics(experiment_id)
        assert len(metrics) >= 13
        assert {metric.split for metric in metrics} >= {
            "train", "validation", "test", "test:k3", "test:k4"
        }
        stages = database.get_experiment_stages(experiment_id)
        assert all(stage.status == "completed" for stage in stages)
        assert sum(stage.stage_name == "model_training" for stage in stages) == 1
        assert any(
            stage.stage_name == "test_evaluation_k3_shared_model"
            for stage in stages
        )
        assert any(
            stage.stage_name == "test_evaluation_k4_shared_model"
            for stage in stages
        )
    assert model_factory_calls == ["lightgbm"]

    batch_config = json.loads(
        (experiment_root / "config" / "batch_config.json").read_text(
            encoding="utf-8"
        )
    )
    assert batch_config["target_sizes"] == [3, 4]
    assert batch_config["model_scope"] == "one model trained on combined sizes"
    sample_sizes = json.loads(
        (experiment_root / "datasets" / "sample_sizes.json").read_text(
            encoding="utf-8"
        )
    )
    assert len(sample_sizes) == 2 * SAMPLE_COUNT
    assert set(sample_sizes.values()) == {3, 4}
    combined_rows = list(
        csv.DictReader(
            (experiment_root / "datasets" / "dataset.csv").open(
                newline="", encoding="utf-8"
            )
        )
    )
    assert len(combined_rows) == 2 * SAMPLE_COUNT
    assert sum(
        int(row["subgraph_size"]) == 3 for row in combined_rows
    ) == SAMPLE_COUNT
    train_rows = list(
        csv.DictReader(
            (experiment_root / "datasets" / "train_dataset.csv").open(
                newline="", encoding="utf-8"
            )
        )
    )
    assert len(train_rows) + 10 + 10 == 2 * SAMPLE_COUNT
    assert (experiment_root / "models" / "model.txt").is_file()
    assert (experiment_root / "predictions" / "predictions.csv").is_file()
    report_path = root / "reports" / "experiments.xlsx"
    assert report_path.is_file()
    workbook = load_workbook(report_path, read_only=True, data_only=True)
    try:
        experiment_rows = list(workbook["Experiments"].values)[1:]
    finally:
        workbook.close()
    assert [row[0] for row in experiment_rows] == [experiment_id]

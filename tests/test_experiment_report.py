import csv
import json

import pytest

from src.experiments.artifact_store import ArtifactStore
from src.experiments.database import ExperimentDatabase
from src.experiments.report import ExperimentReport, SUMMARY_COLUMNS
from openpyxl import load_workbook


def test_update_xlsx_upserts_only_requested_experiment(
    database,
    tmp_path,
):
    path_1 = add_experiment(database, tmp_path, "EXP_001")
    path_2 = add_experiment(database, tmp_path, "EXP_002")

    report_path = tmp_path / "reports" / "experiments.xlsx"
    report = ExperimentReport(database)

    report.update_xlsx(report_path, "EXP_001")
    report.update_xlsx(report_path, "EXP_002")
    report.update_xlsx(report_path, "EXP_002")

    workbook = load_workbook(report_path, read_only=True)
    rows = list(workbook["Experiments"].values)
    workbook.close()

    assert rows[0][0] == "experiment_id"
    assert [row[0] for row in rows[1:]] == ["EXP_001", "EXP_002"]
    assert path_1.is_file()
    assert path_2.is_file()


def test_update_xlsx_writes_stage_artifact_prediction_and_metric_sheets(
    database,
    tmp_path,
):
    add_experiment(database, tmp_path, "EXP_001")
    database.register_stage("EXP_001", "training", "completed", duration_seconds=1.5)
    database.register_artifact(
        "ART_001",
        "EXP_001",
        "models",
        "model.bin",
        "EXP_001/models/model.bin",
        "application/octet-stream",
    )
    database.register_prediction(
        "EXP_001",
        "sample-1",
        0.75,
        actual_value=0.8,
        prediction_error=-0.05,
    )
    database.register_evaluation("EXP_001", "mae", 0.05)

    report_path = tmp_path / "reports" / "experiments.xlsx"
    ExperimentReport(database).update_xlsx(report_path, "EXP_001")

    workbook = load_workbook(report_path, read_only=True)
    assert workbook["Stages"].max_row == 2
    assert workbook["Artifacts"].max_row == 2
    assert workbook["Predictions"].max_row == 2
    assert workbook["Evaluations"].max_row == 2
    workbook.close()


def test_tracking_dataframes_have_explicit_schema_and_nullable_dtypes(
    database,
    tmp_path,
):
    add_experiment(database, tmp_path, "EXP_001")
    add_experiment(database, tmp_path, "EXP_002")
    stage = database.register_stage("EXP_001", "validation", "completed")
    database.register_metric(
        experiment_id="EXP_001",
        stage_id=stage.stage_id,
        metric_name="mae",
        split="validation",
        metric_value=0.125,
    )
    report = ExperimentReport(database)

    frames = report.to_dataframes()

    assert set(frames) == {"experiments", "stages", "metrics"}
    for table, frame in frames.items():
        assert list(frame.columns) == list(report.DATAFRAME_COLUMNS[table])
    assert str(frames["experiments"]["experiment_id"].dtype) == "string"
    assert str(frames["experiments"]["created_at"].dtype) == (
        "datetime64[ns, UTC]"
    )
    assert str(frames["stages"]["stage_id"].dtype) == "Int64"
    assert str(frames["metrics"]["metric_value"].dtype) == "Float64"
    assert frames["metrics"].iloc[0]["stage_id"] == stage.stage_id

    filtered = report.to_dataframes("EXP_001")
    assert filtered["experiments"]["experiment_id"].tolist() == ["EXP_001"]
    assert set(filtered["stages"]["experiment_id"].dropna()) == {"EXP_001"}
    assert set(filtered["metrics"]["experiment_id"].dropna()) == {"EXP_001"}


def test_empty_dataframes_keep_declared_columns_and_types(database):
    frames = ExperimentReport(database).to_dataframes()

    assert all(frame.empty for frame in frames.values())
    assert list(frames["metrics"].columns) == list(
        ExperimentReport.DATAFRAME_COLUMNS["metrics"]
    )
    assert str(frames["metrics"]["stage_id"].dtype) == "Int64"


def test_tracking_dataframe_xlsx_round_trip_preserves_metric_links(
    database,
    tmp_path,
):
    add_experiment(database, tmp_path, "EXP_001")
    stage = database.register_stage("EXP_001", "test_evaluation", "completed")
    database.register_metric(
        experiment_id="EXP_001",
        stage_id=stage.stage_id,
        metric_name="rmse",
        split="test",
        metric_value=0.375,
    )
    report = ExperimentReport(database)
    output_path = tmp_path / "tracking.xlsx"
    experiments_before = database.list_experiments()
    stages_before = database.list_stages()
    metrics_before = database.get_experiment_metrics()

    report.to_dataframes()
    report.export_dataframes_xlsx(output_path)
    frames = ExperimentReport.read_dataframes_xlsx(output_path)

    assert set(frames) == {"experiments", "stages", "metrics"}
    metric = frames["metrics"].iloc[0]
    assert metric["experiment_id"] == "EXP_001"
    assert metric["stage_id"] == stage.stage_id
    assert metric["metric_name"] == "rmse"
    assert metric["split"] == "test"
    assert metric["metric_value"] == pytest.approx(0.375)
    assert str(frames["metrics"]["recorded_at"].dtype) == (
        "datetime64[ns, UTC]"
    )

    with pytest.raises(FileExistsError):
        report.export_dataframes_xlsx(output_path)

    assert database.list_experiments() == experiments_before
    assert database.list_stages() == stages_before
    assert database.get_experiment_metrics() == metrics_before


@pytest.fixture
def database(tmp_path):
    repository = ExperimentDatabase(tmp_path / "registry.sqlite3")
    yield repository
    repository.close()


def add_experiment(
    database,
    tmp_path,
    experiment_id="EXP_001",
    status="COMPLETED",
    configuration=None,
):
    if configuration is None:
        configuration = {
            "graph": {
                "graph_id": "graph-1",
                "path": None,
                "source": None,
            },
            "subgraph_generation": {
                "method": "random_paths",
                "sample_count": 20,
                "path_length": 3,
                "parameters": {},
            },
            "features": {
                "feature_set_id": "structural-v1",
                "feature_names": ["degree_mean", "degree_max"],
                "objective_name": "tds",
            },
            "dataset_split": {
                "strategy": "random",
                "test_ratio": 0.2,
                "validation_size": 0.25,
                "random_state": 11,
                "parameters": {},
            },
            "model": {
                "model_type": "lightgbm",
                "hyperparameters": {
                    "num_leaves": 5,
                    "n_estimators": 10,
                },
            },
        }

    store = ArtifactStore(tmp_path / experiment_id)
    config_path = store.save_json(
        "config",
        "experiment.json",
        configuration,
    )

    database.create_experiment(
        experiment_id=experiment_id,
        name=f"Experiment {experiment_id}",
        description="Report test",
        status=status,
        created_at="2026-09-30T10:00:00Z",
        configuration_path=str(config_path),
    )
    return config_path


def read_report(path):
    with path.open("r", encoding="utf-8", newline="") as csv_file:
        return list(csv.DictReader(csv_file))


def test_generates_csv_report_with_deterministic_headers(
    database,
    tmp_path,
):
    add_experiment(database, tmp_path)
    output_path = tmp_path / "reports" / "summary.csv"

    returned_path = ExperimentReport(database).write_csv(output_path)

    assert returned_path == output_path
    with output_path.open("r", encoding="utf-8", newline="") as csv_file:
        reader = csv.reader(csv_file)
        assert next(reader) == list(SUMMARY_COLUMNS)


def test_one_experiment_produces_one_summary_row(database, tmp_path):
    add_experiment(database, tmp_path)

    rows = read_report(
        ExperimentReport(database).write_csv(tmp_path / "summary.csv")
    )

    assert len(rows) == 1
    assert rows[0]["experiment_id"] == "EXP_001"


def test_multiple_experiments_produce_multiple_rows(database, tmp_path):
    add_experiment(database, tmp_path, "EXP_001")
    add_experiment(database, tmp_path, "EXP_002")

    rows = read_report(
        ExperimentReport(database).write_csv(tmp_path / "summary.csv")
    )

    assert {row["experiment_id"] for row in rows} == {
        "EXP_001",
        "EXP_002",
    }


def test_report_includes_status_model_features_and_split(
    database,
    tmp_path,
):
    add_experiment(database, tmp_path, status="RUNNING")

    [row] = read_report(
        ExperimentReport(database).write_csv(tmp_path / "summary.csv")
    )

    assert row["status"] == "RUNNING"
    assert row["model_type"] == "lightgbm"
    assert json.loads(row["model_hyperparameters_json"]) == {
        "n_estimators": 10,
        "num_leaves": 5,
    }
    assert row["feature_set_id"] == "structural-v1"
    assert row["number_of_features"] == "2"
    assert json.loads(row["feature_names_json"]) == [
        "degree_mean",
        "degree_max",
    ]
    assert row["dataset_split_strategy"] == "random"
    assert row["validation_size"] == "0.25"
    assert row["test_ratio"] == "0.2"
    assert row["random_state"] == "11"


def test_report_includes_stage_count_evaluation_and_prediction_count(
    database,
    tmp_path,
):
    add_experiment(database, tmp_path)
    database.register_stage(
        "EXP_001",
        "training",
        "completed",
        duration_seconds=2.5,
    )
    database.register_stage(
        "EXP_001",
        "evaluation",
        "completed",
        duration_seconds=0.5,
    )
    database.register_prediction("EXP_001", "sample-1", 0.3)
    database.register_prediction("EXP_001", "sample-2", 0.7)
    database.register_evaluation("EXP_001", "mae", 0.125)
    database.register_evaluation("EXP_001", "rmse", 0.25)

    [row] = read_report(
        ExperimentReport(database).write_csv(tmp_path / "summary.csv")
    )

    assert row["recorded_stage_count"] == "2"
    assert row["prediction_count"] == "2"
    assert json.loads(row["evaluation_metrics_json"]) == {
        "mae": 0.125,
        "rmse": 0.25,
    }


def test_filtering_by_one_or_many_experiment_ids(database, tmp_path):
    add_experiment(database, tmp_path, "EXP_001")
    add_experiment(database, tmp_path, "EXP_002")
    add_experiment(database, tmp_path, "EXP_003")
    report = ExperimentReport(database)

    one_row_path = report.write_csv(
        tmp_path / "one.csv",
        experiment_ids="EXP_002",
    )
    selected_path = report.write_csv(
        tmp_path / "selected.csv",
        experiment_ids=["EXP_003", "EXP_001"],
    )

    assert [row["experiment_id"] for row in read_report(one_row_path)] == [
        "EXP_002",
    ]
    assert [row["experiment_id"] for row in read_report(selected_path)] == [
        "EXP_003",
        "EXP_001",
    ]


def test_empty_database_produces_csv_with_headers(database, tmp_path):
    output_path = ExperimentReport(database).write_csv(
        tmp_path / "empty.csv"
    )

    with output_path.open("r", encoding="utf-8", newline="") as csv_file:
        rows = list(csv.reader(csv_file))

    assert rows == [list(SUMMARY_COLUMNS)]


def test_nested_configuration_is_serialized_deterministically(
    database,
    tmp_path,
):
    configuration = {
        "graph": {
            "graph_id": "graph-1",
            "parameters": {"z": 1, "a": 2},
        },
        "subgraph_generation": {
            "method": "paths",
            "sample_count": 4,
            "parameters": {"outer": {"z": 3, "a": 1}},
        },
        "features": {
            "feature_set_id": "set-v1",
            "feature_names": ["degree"],
            "objective_name": "tds",
        },
        "dataset_split": {
            "strategy": "random",
            "test_ratio": 0.2,
            "validation_size": 0.25,
            "random_state": 11,
            "parameters": {"shuffle": True},
        },
        "model": {
            "model_type": "example",
            "hyperparameters": {"z": 2, "a": {"y": 1, "b": 0}},
        },
    }
    add_experiment(
        database,
        tmp_path,
        configuration=configuration,
    )

    first_path = ExperimentReport(database).write_csv(
        tmp_path / "first.csv"
    )
    second_path = ExperimentReport(database).write_csv(
        tmp_path / "second.csv"
    )

    [first_row] = read_report(first_path)
    [second_row] = read_report(second_path)

    assert (
        first_row["model_hyperparameters_json"]
        == second_row["model_hyperparameters_json"]
    )
    assert first_row["model_hyperparameters_json"] == (
        '{"a":{"b":0,"y":1},"z":2}'
    )


def test_report_does_not_modify_experiment_records(database, tmp_path):
    add_experiment(database, tmp_path)
    database.register_stage(
        "EXP_001",
        "training",
        "completed",
        duration_seconds=1.0,
    )

    experiment_before = database.get_experiment("EXP_001")
    stages_before = database.get_experiment_stages("EXP_001")
    artifacts_before = database.get_experiment_artifacts("EXP_001")

    ExperimentReport(database).write_csv(tmp_path / "summary.csv")
    ExperimentReport(database).get_experiment_details("EXP_001")

    assert database.get_experiment("EXP_001") == experiment_before
    assert database.get_experiment_stages("EXP_001") == stages_before
    assert database.get_experiment_artifacts("EXP_001") == artifacts_before


def test_detailed_report_includes_database_records(database, tmp_path):
    add_experiment(database, tmp_path)
    database.register_stage("EXP_001", "training", "completed")
    database.register_artifact(
        "ART_001",
        "EXP_001",
        "logs",
        "run.log",
        "EXP_001/logs/run.log",
        "text/plain",
    )
    database.register_prediction("EXP_001", "sample-1", 0.5)
    database.register_evaluation("EXP_001", "mae", 0.1)

    details = ExperimentReport(database).get_experiment_details("EXP_001")

    assert len(details["stages"]) == 1
    assert len(details["artifacts"]) == 1
    assert len(details["predictions"]) == 1
    assert len(details["evaluations"]) == 1